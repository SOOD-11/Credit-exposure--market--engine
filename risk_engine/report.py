"""Excel report (live formulas) and backtest charts."""
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from .bs import bs_greeks
from .config import FIG_DIR, RISK_FREE_RATE, PARAMS
from .exposure import mark_to_market
from .valuation import trailing_vol

ARIAL = "Arial"
BLUE, GREEN = "0000FF", "008000"
HDR_FILL = PatternFill("solid", start_color="1F3864")
MONEY = '$#,##0;($#,##0);-'


def _f(bold=False, color="000000", italic=False, size=10):
    return Font(name=ARIAL, size=size, bold=bold, color=color, italic=italic)


def commentary(name, pos, prices, idx, mtm_today, im, collateral, pot_exp, limit):
    """One-sentence plain-English note on the day's move and its main driver."""
    prev = mark_to_market(pos, prices, idx - 1)
    delta = mtm_today.set_index(["ticker", "instrument"]).mtm - prev.set_index(["ticker", "instrument"]).mtm
    by_t = delta.groupby(level=0).sum()
    top = by_t.abs().idxmax()
    total = delta.sum()
    state = "client is in the money, so no uncovered exposure today" if mtm_today.mtm.sum() >= 0 else "client owes the bank on a net basis"
    status = "ABOVE limit" if pot_exp > limit else "within limit"
    return (f"Client MTM {'rose' if total >= 0 else 'fell'} ${abs(total):,.0f} vs prior close, mainly {top} "
            f"({by_t[top]:+,.0f}); {state}. Potential exposure after a 99%/10d shock ${pot_exp:,.0f} is {status} "
            f"(${limit:,.0f}); margin ${im:,.0f} vs collateral ${collateral:,.0f}.")


def build_workbook(path, asof, clients, trades, prices, idx, margins, backtests, params=PARAMS):
    """margins: {client_id: initial_margin() dict}; backtests: {client_id: summarise() dict}."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Counterparty Exposure & Initial Margin Report (SYNTHETIC CLIENTS)"
    ws["A1"].font = _f(True, size=14)
    ws["A2"] = f"As-of date: {asof:%d %b %Y}"
    ws["A2"].font = _f()
    ws["A3"] = ("Blue = hardcoded input or model output from Python; black = Excel formula; green = link to another sheet. "
                "Exposure = MAX(-MTM - collateral, 0). Margin call = MAX(initial margin - collateral, 0). "
                "Potential exposure = MAX(-(MTM - IM) - collateral, 0), the figure the limit is tested against.")
    ws["A3"].font = _f(italic=True, size=9)
    ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A3:M3")
    ws.row_dimensions[3].height = 42
    heads = ["Client", "Gross Notional ($)", "MTM to Client ($)", "Collateral Held ($)", "Current Exposure ($)",
             "Initial Margin 99%/10d ($)", "Expected Shortfall ($)", "Margin Call ($)", "Potential Exposure ($)",
             "Exposure Limit ($)", "Limit Status", "Margin Status", "Daily Commentary"]
    for c, h in enumerate(heads, 1):
        cell = ws.cell(row=6, column=c, value=h)
        cell.font, cell.fill = _f(True, "FFFFFF"), HDR_FILL
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws.row_dimensions[6].height = 32

    r = 7
    for cl in clients.itertuples():
        pos = trades[trades.client_id == cl.client_id].reset_index(drop=True)
        m = mark_to_market(pos, prices, idx)
        vol = trailing_vol(prices, idx)
        # ----- client sheet -----
        cs = wb.create_sheet(f"Client{cl.client_id}")
        cs["A1"] = f"Client {cl.client_id}: {cl.name} (synthetic)"
        cs["A1"].font = _f(True, size=12)
        cs["A2"] = ("Qty in shares (+ long swap, - written put). MTM = Qty x (Current unit value - Entry value). "
                    "Unit values/Greeks are Black-Scholes outputs from Python (blue).")
        cs["A2"].font = _f(italic=True, size=9)
        hdr = ["Ticker", "Instrument", "Qty", "Strike", "Expiry", "Entry Value", "Current Unit Value", "MTM ($)",
               "Notional ($)", "Delta", "Gamma", "Vega (per vol pt)"]
        for c, h in enumerate(hdr, 1):
            cell = cs.cell(row=4, column=c, value=h)
            cell.font, cell.fill = _f(True, "FFFFFF"), HDR_FILL
            cell.alignment = Alignment(wrap_text=True, horizontal="center")
        row = 5
        for p in m.itertuples():
            cs.cell(row=row, column=1, value=p.ticker)
            cs.cell(row=row, column=2, value=p.instrument)
            cs.cell(row=row, column=3, value=float(p.qty)).number_format = '#,##0;(#,##0)'
            if p.instrument == "PUT":
                cs.cell(row=row, column=4, value=float(p.strike)).number_format = '0.00'
                cs.cell(row=row, column=5, value=pd.Timestamp(p.expiry).to_pydatetime()).number_format = 'dd-mmm-yy'
                T = max((pd.Timestamp(p.expiry) - asof).days / 365.0, 1e-6)
                g = bs_greeks(p.spot, p.strike, T, RISK_FREE_RATE, vol[p.ticker], "put")
                cs.cell(row=row, column=10, value=float(g["delta"])).number_format = '0.000'
                cs.cell(row=row, column=11, value=float(g["gamma"])).number_format = '0.0000'
                cs.cell(row=row, column=12, value=float(g["vega"]) / 100).number_format = '0.000'
            else:
                cs.cell(row=row, column=10, value=1.0).number_format = '0.000'
            cs.cell(row=row, column=6, value=float(p.entry_value)).number_format = '0.00'
            cs.cell(row=row, column=7, value=float(p.unit_value)).number_format = '0.00'
            cs.cell(row=row, column=8, value=f"=C{row}*(G{row}-F{row})").number_format = MONEY
            cs.cell(row=row, column=9, value=float(p.notional)).number_format = MONEY
            for c in range(1, 13):
                cell = cs.cell(row=row, column=c)
                is_formula = c == 8
                cell.font = _f(color="000000" if is_formula or c in (1, 2) else BLUE)
            row += 1
        last = row - 1
        cs.cell(row=row, column=1, value="Total").font = _f(True)
        for c, letter in [(8, "H"), (9, "I")]:
            cell = cs.cell(row=row, column=c, value=f"=SUM({letter}5:{letter}{last})")
            cell.font, cell.number_format = _f(True), MONEY
        total_row = row
        for c, w in zip("ABCDEFGHIJKL", [9, 11, 12, 9, 11, 12, 15, 15, 15, 8, 9, 15]):
            cs.column_dimensions[c].width = w
        cs.freeze_panes = "A5"

        # ----- summary row -----
        coll = float(cl.collateral)
        im, es = margins[cl.client_id]["im"], margins[cl.client_id]["es"]
        mtm_total = float(m.mtm.sum())
        pot = max(-(mtm_total - im) - coll, 0.0)
        ws.cell(row=r, column=1, value=f"{cl.client_id}: {cl.name}")
        ws.cell(row=r, column=2, value=float(cl.gross_notional))
        ws.cell(row=r, column=3, value=f"='Client{cl.client_id}'!H{total_row}")
        ws.cell(row=r, column=4, value=coll)
        ws.cell(row=r, column=5, value=f"=MAX(-C{r}-D{r},0)")
        ws.cell(row=r, column=6, value=float(im))
        ws.cell(row=r, column=7, value=float(es))
        ws.cell(row=r, column=8, value=f"=MAX(F{r}-D{r},0)")
        ws.cell(row=r, column=9, value=f"=MAX(-(C{r}-F{r})-D{r},0)")
        ws.cell(row=r, column=10, value=float(cl.exposure_limit))
        ws.cell(row=r, column=11, value=f'=IF(I{r}>J{r},"BREACH","OK")')
        ws.cell(row=r, column=12, value=f'=IF(H{r}>0,"CALL","OK")')
        ws.cell(row=r, column=13, value=commentary(cl.name, pos, prices, idx, m, im, coll, pot, float(cl.exposure_limit)))
        colour = {1: "000000", 2: BLUE, 3: GREEN, 4: BLUE, 5: "000000", 6: BLUE, 7: BLUE, 8: "000000", 9: "000000",
                  10: BLUE, 11: "000000", 12: "000000", 13: BLUE}
        for c in range(1, 14):
            cell = ws.cell(row=r, column=c)
            cell.font = _f(color=colour[c])
            if 2 <= c <= 10:
                cell.number_format = MONEY
            if c == 13:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 70
        r += 1
    for c, w in zip("ABCDEFGHIJKLM", [32, 16, 16, 16, 16, 18, 16, 16, 18, 16, 12, 12, 80]):
        ws.column_dimensions[c].width = w
    ws.freeze_panes = "B7"

    # ----- backtest sheet -----
    bs = wb.create_sheet("Backtest")
    bs["A1"] = "Initial-margin model backtest (model outputs from Python)"
    bs["A1"].font = _f(True, size=12)
    bs["A2"] = ("An 'exception' = realised 10-day loss larger than the margin. A 99% model should show about 1%. "
                "Kupiec p-value < 0.05 means the exception rate is statistically inconsistent with 1%.")
    bs["A2"].font = _f(italic=True, size=9)
    heads = ["Client", "Sample", "Test days", "Exceptions", "Exception rate", "Expected rate", "Kupiec LR", "Kupiec p-value"]
    for c, h in enumerate(heads, 1):
        cell = bs.cell(row=4, column=c, value=h)
        cell.font, cell.fill = _f(True, "FFFFFF"), HDR_FILL
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
    row = 5
    for cid, s in backtests.items():
        for key, label in [("daily_overlapping", "Every day (overlapping)"), ("non_overlapping", "Every 10th day (independent)")]:
            d = s[key]
            vals = [cid, label, d["n"], d["exceptions"], f"=D{row}/C{row}", d["expected_rate"], d["kupiec_lr"], d["kupiec_p"]]
            for c, v in enumerate(vals, 1):
                cell = bs.cell(row=row, column=c, value=v)
                cell.font = _f(color="000000" if c in (1, 2, 5) else BLUE)
            bs.cell(row=row, column=5).number_format = '0.00%'
            bs.cell(row=row, column=6).number_format = '0.00%'
            bs.cell(row=row, column=7).number_format = '0.00'
            bs.cell(row=row, column=8).number_format = '0.000'
            row += 1
    for c, w in zip("ABCDEFGH", [8, 30, 11, 12, 14, 14, 12, 14]):
        bs.column_dimensions[c].width = w

    # ----- assumptions sheet -----
    asm = wb.create_sheet("Assumptions")
    lines = [
        ("Clients and trades", "Synthetic (seeded random entry dates). No real client data."),
        ("Prices", "Daily adjusted closes, 20 US large caps, Jan 2015 - Dec 2025 (from the earlier yfinance download). No FX or rates."),
        ("Instruments", "Total-return equity swaps (client long) and written European puts (Black-Scholes, no dividends)."),
        ("Risk-free rate", f"{RISK_FREE_RATE:.1%} flat (assumption, not looked up)."),
        ("Option volatility", "Trailing 252-day realised volatility; held constant across scenarios (no vega risk in scenarios)."),
        ("Margin method", f"Historical simulation: overlapping {params.horizon}-day returns over a {params.lookback}-day lookback applied to today's positions; {params.confidence:.0%} quantile."),
        ("Concentration", f"Names above {params.conc_threshold:.0%} of gross notional get scenario moves scaled by sqrt(1.5) (longer close-out)."),
        ("Collateral", "Flat 10% of gross notional (the 'rule of thumb' the model is compared against)."),
        ("Netting", "One netting set per client; no cross-client offsets; no thresholds or minimum transfer amounts."),
        ("Not modelled", "Intraday moves, liquidity costs, wrong-way correlation beyond what is in the data, credit quality of the client."),
    ]
    asm["A1"] = "Assumptions & limitations"
    asm["A1"].font = _f(True, size=12)
    for i, (k, v) in enumerate(lines, 3):
        asm.cell(row=i, column=1, value=k).font = _f(True)
        asm.cell(row=i, column=2, value=v).font = _f()
        asm.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    asm.column_dimensions["A"].width = 22
    asm.column_dimensions["B"].width = 120
    wb.save(path)


def plot_backtest(cid, name, bt, path):
    fig, ax = plt.subplots(figsize=(11, 4.2))
    ax.plot(bt.index, bt.margin / 1e6, color="#1F3864", lw=1.2, label="Initial margin (99%, 10d)")
    ax.fill_between(bt.index, 0, bt.realised_loss.clip(lower=0) / 1e6, color="#9DB7D5", alpha=0.6, label="Realised 10-day loss")
    exc = bt[bt.exception]
    ax.scatter(exc.index, exc.realised_loss / 1e6, color="#C00000", s=14, zorder=3, label=f"Exceptions ({len(exc)})")
    ax.set_ylabel("$ millions")
    ax.set_title(f"Client {cid}: {name} - margin vs realised loss (synthetic)")
    ax.legend(loc="upper left", fontsize=8)
    ax.grid(alpha=0.25)
    fig.tight_layout()
    fig.savefig(path, dpi=140)
    plt.close(fig)
